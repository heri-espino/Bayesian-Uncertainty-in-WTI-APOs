"""Lightweight source checks for the journal-specific REMEF Word draft."""
from __future__ import annotations

import importlib.util
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "remef" / "build.py"


def _build_module():
    spec = importlib.util.spec_from_file_location("_remef_build", SOURCE)
    assert spec is not None and spec.loader is not None
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_rem_e_f_inputs_and_blinding():
    b = _build_module()
    b.validate()
    tex = b.make_latex(with_author=False)
    assert "Heriberto Espino-Montelongo" not in tex
    assert "github.com/heri-espino" not in tex
    assert r"\citet{matiasgonzalez2019}" in tex
    assert r"\citet{ortizjimenez2026}" in tex
    assert r"\citet{gildertsiaras2020}" in tex
    assert r"\section{Results}" in tex
    assert r"\section{Conclusion}" in tex
    assert r"\begin{tabular}" in tex
    assert tex.count(r"\includegraphics") >= 2
    assert r"\begin{table*}" not in tex
    assert r"\begin{figure*}" not in tex
    assert r"\eqref{" not in tex
    assert r"\ref{" not in tex
    assert r"\label{" not in tex


def test_rem_e_f_author_and_bilingual_abstract_limits():
    b = _build_module()
    author = b.make_latex(with_author=True)
    assert "Heriberto Espino-Montelongo" in author
    fm = (ROOT / "paper" / "remef" / "frontmatter.tex").read_text(encoding="utf-8")
    for start, end in (
        (r"\section*{Abstract}", r"\textbf{Keywords:}"),
        (r"\section*{Resumen}", r"\textbf{Palabras clave:}"),
    ):
        segment = fm.split(start, 1)[1].split(end, 1)[0]
        segment = re.sub(r"\\textbf\{([^}]*)\}", r"\1", segment)
        words = segment.split()
        assert len(words) <= 180, (start, len(words))
    assert len("Posterior Integration and Volatility Inputs in WTI Asian Options".split()) <= 10
    assert len("Integración posterior y volatilidad en opciones asiáticas WTI".split()) <= 10


def test_rem_e_f_sources_distinct_from_original_wiley():
    b = _build_module()
    tex = b.make_latex(with_author=False)
    assert "Posterior Integration and Volatility Inputs in WTI Asian Options" in tex
    assert r"\documentclass[HARVARD,Utopia2COL]" not in tex
    assert r"\subsection{AI-assisted manuscript preparation}" in tex
    assert r"\noindent\textbf{Table " in tex
    assert r"\noindent\textbf{Figure " in tex

def test_remef_apa_and_required_section_postprocessor_exist():
    """Keep APA and REMEF section structure explicit in the reproducible build."""
    style = ROOT / "paper" / "remef" / "apa.csl"
    post = ROOT / "paper" / "remef" / "postprocess.py"
    builder = SOURCE.read_text(encoding="utf-8")
    assert style.is_file()
    assert "American Psychological Association" in style.read_text(encoding="utf-8") or "APA" in style.read_text(encoding="utf-8")
    assert post.is_file()
    content = post.read_text(encoding="utf-8")
    for heading in ("Literature review", "Methodology", "Results and discussion", "Conclusions and recommendations"):
        assert heading in content
    assert '"--csl=" + str(HERE / "apa.csl")' in builder
    assert "prepare_docx(path)" in builder

def test_confirmed_funding_and_coi_in_both_versions():
    b = _build_module()
    private = b.make_latex(with_author=True)
    reviewer = b.make_latex(with_author=False)
    for content in (private, reviewer):
        assert "This research received no external funding." in content
        assert "The author declares no conflicts of interest." in content
        assert "To be completed and verified" not in content
    for sensitive in (
        "Heriberto Espino-Montelongo",
        "heriberto.espinomo@udlap.mx",
        "0009-0009-1230-2931",
        "heri-espino.github.io",
        "linkedin.com/in/heri-espino",
    ):
        assert sensitive in private
        assert sensitive not in reviewer
    assert r"\footnote{" in private
    assert "Ex Hacienda Santa Catarina" in private
    assert "San Andrés" not in reviewer
    assert "+52 222 810 1202" in private
    assert "+52 222 810 1202" not in reviewer
    assert "[author to supply" not in private


def test_actual_pandoc_heading_text_is_supported(tmp_path):
    """Regression: Pandoc creates unnumbered headings in Word (run #39158592891).

    Exercise actual python-docx postprocessing on minimal Pandoc-shaped source,
    including moves between sections, so this bug fails lightweight pytest CI.
    """
    from docx import Document
    spec = importlib.util.spec_from_file_location(
        "_remef_postprocess", ROOT / "paper" / "remef" / "postprocess.py"
    )
    assert spec is not None and spec.loader is not None
    processor = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(processor)

    doc = Document()
    doc.add_paragraph("Posterior Integration and Volatility Inputs in WTI Asian Options")
    doc.add_paragraph("Integración posterior y volatilidad en opciones asiáticas WTI")
    doc.add_paragraph("Abstract", style="Heading 1")
    doc.add_paragraph("Keywords: Bayesian option pricing")
    doc.add_paragraph("Resumen", style="Heading 1")
    doc.add_paragraph("JEL: C11; C15; C58; G13; G17; Q40.")
    doc.add_paragraph("Introduction", style="Heading 1")
    doc.add_paragraph("Volatility is not observed directly")
    doc.add_paragraph("Asian options depend on a price average")
    doc.add_paragraph("Bayesian option pricing is also well established")
    doc.add_paragraph("WTI APOs provide a different test setting")
    doc.add_paragraph("The distinction between historical and option-implied information")
    doc.add_paragraph("Posterior uncertainty and Asian-option pricing", style="Heading 1")
    for title in (
        "Physical inference and risk-neutral valuation",
        "Posterior integration and price uncertainty",
        "Synthetic mechanism experiment",
    ):
        doc.add_paragraph(title, style="Heading 2")
        doc.add_paragraph("A reproducible scientific calculation.")
    doc.add_paragraph("Data and empirical design", style="Heading 1")
    for title in (
        "WTI Average Price Options",
        "Out-of-sample volatility comparisons",
        "Cross-instrument validation",
        "Independent extended Databento settlement-mark benchmark",
        "AI-assisted manuscript preparation",
    ):
        doc.add_paragraph(title, style="Heading 2")
        doc.add_paragraph("Data and analysis use earlier observations.")
    doc.add_paragraph("Results", style="Heading 1")
    for title in (
        "When does posterior integration matter?",
        "Posterior integration in the WTI sample",
        "Historical versus option-implied volatility",
        "Cross-instrument validation using vanilla-WTI options",
        "Extended conditional settlement-mark evidence",
    ):
        doc.add_paragraph(title, style="Heading 2")
        doc.add_paragraph("The results are calculated from saved experiments.")
    doc.add_paragraph("Robustness and interpretation", style="Heading 1")
    doc.add_paragraph("Conclusion", style="Heading 1")
    doc.add_paragraph("Data availability", style="Heading 1")
    doc.add_paragraph("External data subject to licensing.")
    doc.add_paragraph("Funding", style="Heading 1")
    doc.add_paragraph("This research received no external funding.")
    doc.add_paragraph("Conflicts of interest", style="Heading 1")
    doc.add_paragraph("The author declares no conflicts of interest.")
    doc.add_paragraph("References", style="Heading 1")
    for rows, cols in ((5, 3), (5, 5), (4, 5), (4, 3)):
        tab = doc.add_table(rows=rows, cols=cols)
        for cell in tab.rows[0].cells:
            cell.text = "Value"
    source = tmp_path / "pandoc_unumbered.docx"
    doc.save(source)
    processor.prepare_docx(source)
    result = Document(source)
    headings = [p.text for p in result.paragraphs if p.style.name.startswith("Heading")]
    for required in (
        "1. Introduction", "2. Literature review", "3. Methodology",
        "3.8 Computational implementation", "3.9 AI-assisted manuscript preparation",
        "4. Results and discussion", "4.5 Independent Databento settlement-mark results",
        "4.6 Robustness, interpretation and limitations",
        "5. Conclusions and recommendations",
    ):
        assert required in headings
    assert "Data and empirical design" not in headings
    assert len(result.tables) == 4
    assert result.tables[0].cell(0, 1).text == "N"
    assert any("The REMEF studies examine" in p.text for p in result.paragraphs)
    assert any("no external funding" in p.text for p in result.paragraphs)
    # Reference-PDF title alignment and REMEF pagination are deterministic.
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    titles = [
        "Posterior Integration and Volatility Inputs in WTI Asian Options",
        "Integración posterior y volatilidad en opciones asiáticas WTI",
    ]
    for title in titles:
        para = next(p for p in result.paragraphs if p.text.strip() == title)
        assert para.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert result.sections[0].footer._element.xpath(".//w:fldSimple")
    assert "PAGE" in result.sections[0].footer._element.xml



def test_remef_manual_workflow_produces_math_reference_not_broken_lo_pdf():
    """Guard against regressing to LibreOffice-only rendering of Word OMML."""
    path = ROOT / ".github" / "workflows" / "build-remef.yml"
    workflow = path.read_text(encoding="utf-8")
    assert "workflow_dispatch:" in workflow
    assert "on:\\n  push:" not in workflow
    assert "ubuntu-24.04" in workflow
    for action in ("actions/checkout@v5", "actions/setup-python@v6",
                   "actions/upload-artifact@v6"):
        assert action in workflow
    assert "--from=docx --to=pdf --pdf-engine=xelatex" in workflow
    assert "remef_math_reference.pdf" in workflow
    assert "unicodedata.normalize(" in workflow
    assert "lo_layout" in workflow
    assert "Greek math symbols are absent from PDF preview" not in workflow
    assert "            paper/remef/build/remef_blinded_draft.pdf" not in workflow



def test_word_builder_prefers_native_svg_over_png():
    b = _build_module()
    src = b.make_latex(with_author=False)
    assert "fig01_mechanism_map.svg" in src
    assert "fig03_forward_q_cluster_bootstrap.svg" in src
    assert "fig01_mechanism_map.png" not in src
    assert "fig03_forward_q_cluster_bootstrap.png" not in src
    assert "vector_figures" in src
    workflow = (ROOT / ".github" / "workflows" / "build-remef.yml").read_text(encoding="utf-8")
    assert "inkscape librsvg2-bin" in workflow
    assert 'native_svg = xml.count("svgBlip")' in workflow
    assert "paper/remef/build/vector_figures/*.svg" in workflow


def test_figure3_vector_title_and_sanitization(tmp_path):
    import xml.etree.ElementTree as ET
    b = _build_module()
    p = tmp_path / "fig03_forward_q_cluster_bootstrap.svg"
    p.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 450 230">'
        '<metadata>Old PDF metadata</metadata>'
        '<path d="M0 0 L5 5"/></svg>', encoding="utf-8",
    )
    b._sanitize_and_repair_svg(p, repair_figure_3=True)
    root = ET.parse(p).getroot()
    names = [el.tag.split("}")[-1] for el in root.iter()]
    assert "metadata" not in names
    assert "rect" in names and "text" in names
    labels = [el.text for el in root.iter() if el.tag.endswith("}text")]
    assert labels == ["Forward-in-time volatility comparisons across expiries"]
    assert any(el.tag.endswith("}path") for el in root.iter())


def test_remef_reconciled_databento_sections():
    """Journal Word build reads REMEF-only revised sections, never rewrites Wiley source."""
    b = _build_module()
    assert all("paper/remef/sections" in p.as_posix() for p in b.BODY_SECTIONS)
    tex = b.make_latex(with_author=False)
    assert "Independent extended Databento settlement-mark benchmark" in tex
    assert "Extended conditional settlement-mark evidence" in tex
    assert "October 2026 expiry" in tex
    assert "97,341" in tex and "111,585" in tex
    assert "0.2089" in tex and "0.2098" in tex
    assert "2026" in tex
    assert "Data availability" in tex
    assert tex.count(r"\noindent\textbf{Table ") >= 5
    assert "date-balanced" in tex or "date-balanced" in (
        ROOT / "paper" / "remef" / "sections" / "03_data_design.tex"
    ).read_text(encoding="utf-8")

def test_remef_workflow_checks_current_six_native_tables():
    """Table 2 has two independently editable Word tables (Panels A and B)."""
    b = _build_module()
    source = "\n".join(p.read_text(encoding="utf-8") for p in b.BODY_SECTIONS)
    assert source.count(r"\begin{tabular}") == 6, (
        "Review the Word build if the manuscript table count changes"
    )
    workflow = (ROOT / ".github/workflows/build-remef.yml").read_text(encoding="utf-8")
    assert "tables != 6" in workflow
    assert "len(doc.tables) != 6" in workflow


def test_remef_final_references_and_caveats():
    """Protect the close-prior-work citation and the evidence-as-of boundary."""
    b = _build_module()
    text = b.make_latex(with_author=False)
    bib = (ROOT / "paper/remef/references.bib").read_text(encoding="utf-8")
    assert r"\citet{romboutsstentoft2014}" in text
    assert r"\citet{sagorbooneghanam2026}" in text
    assert "@article{sagorbooneghanam2026" in bib
    assert "2,319 option-date rows" in text
    assert "2,366 same-contract" in text
    assert "first-publication" in text
