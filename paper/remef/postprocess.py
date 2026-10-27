"""Journal-specific, non-destructive layout adaptation of Pandoc Word output.

Keeps native OMML equations editable. Do not imitate publisher-produced DOI/logo headers.
Source: REMEF author instructions and a published article supplied for layout reference.
"""
from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def find(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise ValueError("Expected REMEF paragraph not found: " + prefix)


def insert_before(anchor, text, style=None):
    p = anchor.insert_paragraph_before(text)
    if style:
        other = next(x for x in anchor._parent.paragraphs if x.style.name == style)
        p.style = other.style
    return p


def remove(p):
    element = p._element
    element.getparent().remove(element)



def repair_equation_references(doc):
    """Repair equation labels lost when converting LaTeX math to editable Word math."""
    substitutions = (
        ("Equation\u00a0(4)", "the arithmetic-average definition above"),
        ("Equation (4)", "the arithmetic-average definition above"),
        ("the two ingredients in Equation\u00a0(7)",
         "posterior variance and price curvature in the second-order approximation"),
        ("the two ingredients in Equation (7)",
         "posterior variance and price curvature in the second-order approximation"),
        ("the corresponding second-order approximation in Equation\u00a0(7)",
         "the corresponding second-order approximation"),
        ("the corresponding second-order approximation in Equation (7)",
         "the corresponding second-order approximation"),
    )
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for old, replacement in substitutions:
                if old in run.text:
                    run.text = run.text.replace(old, replacement)

def adapt_sections(doc):
    """Five REMEF-style scientific sections; preserve scientific paragraphs and formulas."""
    # Pandoc writes unnumbered heading text by default; numbering belongs
    # to the REMEF Word layout, not to the LaTeX-to-DOCX converter.
    intro = find(doc, "Introduction")
    if intro.style.name != "Heading 1":
        raise ValueError("Pandoc introduction did not retain Heading 1 style")
    intro.text = "1. Introduction"
    original = find(doc, "Posterior uncertainty and Asian-option pricing")
    insert_before(original, "2. Literature review", "Heading 1")
    for prefix in (
        "Asian options depend on a price average",
        "WTI APOs provide a different test setting",
        "The distinction between historical and option-implied information",
    ):
        original._p.addprevious(find(doc, prefix)._p)
    insert_before(
        original,
        "The REMEF studies examine Asian-option valuation with stochastic-control or "
        "mean-reverting volatility specifications; the crude-oil option study addresses "
        "forecasts of subsequent realized volatility. They are complementary but do not "
        "directly test our question. Here the pricing dynamics are held fixed while we "
        "separate the posterior-integration correction, the dispersion of conditional "
        "model values, and the contribution of previously available volatility inputs "
        "to out-of-sample APO pricing errors.",
    )
    # These are the actual Pandoc-generated Heading 1/2 paragraph texts.
    # Do not assume the LaTeX section numbers will survive conversion.
    new = {
        "Posterior uncertainty and Asian-option pricing": "3. Methodology",
        "Physical inference and risk-neutral valuation": "3.1 Physical-measure inference and risk-neutral valuation",
        "Posterior integration and price uncertainty": "3.2 Posterior integration and price uncertainty",
        "Synthetic mechanism experiment": "3.3 Synthetic experiment",
        "WTI Average Price Options": "3.4 WTI Average Price Options and market data",
        "Out-of-sample volatility comparisons": "3.5 Forward-in-time volatility comparisons",
        "Cross-instrument validation": "3.6 Cross-instrument validation",
        "Independent extended Databento settlement-mark benchmark": "3.7 Extended conditional AAO settlement-mark study",
        "AI-assisted manuscript preparation": "3.9 AI-assisted manuscript preparation",
        "Results": "4. Results and discussion",
        "When does posterior integration matter?": "4.1 When does posterior integration matter?",
        "Posterior integration in the WTI sample": "4.2 Posterior integration in the WTI sample",
        "Historical versus option-implied volatility": "4.3 Historical versus option-implied volatility",
        "Cross-instrument validation using vanilla-WTI options": "4.4 Cross-instrument validation using vanilla-WTI options",
        "Extended conditional settlement-mark evidence": "4.5 Independent Databento settlement-mark results",
        "Robustness and interpretation": "4.6 Robustness, interpretation and limitations",
        "Conclusion": "5. Conclusions and recommendations",
    }
    for p in doc.paragraphs:
        key = p.text.strip()
        if key in new:
            p.text = new[key]
            if key == "Robustness and interpretation":
                p.style = find(doc, "3.1 Physical-measure").style
    remove(find(doc, "Data and empirical design"))
    expected = [
        "1. Introduction", "2. Literature review", "3. Methodology",
        "3.1 Physical-measure inference and risk-neutral valuation",
        "3.2 Posterior integration and price uncertainty",
        "3.3 Synthetic experiment", "3.4 WTI Average Price Options and market data",
        "3.5 Forward-in-time volatility comparisons", "3.6 Cross-instrument validation",
        "3.9 AI-assisted manuscript preparation", "4. Results and discussion",
        "4.1 When does posterior integration matter?",
        "4.2 Posterior integration in the WTI sample",
        "4.3 Historical versus option-implied volatility",
        "4.4 Cross-instrument validation using vanilla-WTI options",
        "4.5 Independent Databento settlement-mark results",
        "4.6 Robustness, interpretation and limitations",
        "5. Conclusions and recommendations",
    ]
    titles = [p.text.strip() for p in doc.paragraphs if p.style.name.startswith("Heading")]
    missing = [title for title in expected if title not in titles]
    if missing:
        raise ValueError(f"REMEF section conversion missed headings: {missing}")
    ai = find(doc, "3.9 AI-assisted manuscript preparation")
    insert_before(ai, "3.8 Computational implementation", "Heading 2")
    insert_before(
        ai,
        "The computations are implemented in Python using NumPy, SciPy and pandas "
        "for inference, simulation and market-data processing, and Matplotlib "
        "for figures. A public source-only repository contains the pricing "
        "library, experiment scripts, a synthetic demonstration, numerical "
        "tests, and source-provenance records. The original vendor observations "
        "and internal market-dependent outputs are retained privately. The "
        "synthetic calculation integrates physical drift analytically "
        "conditional on volatility. For empirical validation, estimated "
        "volatility inputs follow the documented time cutoffs, and errors "
        "are resampled by valuation date to retain within-date dependence.",
    )
    conclusion = find(doc, "5. Conclusions and recommendations")
    insert_before(
        conclusion,
        "Compared with the related Asian-option valuation studies, these results do "
        "not establish that constant volatility is the best structural pricing "
        "model. They show how posterior integration and the choice of volatility "
        "input contribute differently to the pricing errors of one fixed benchmark.",
    )
    insert_before(
        find(doc, "Data availability"),
        "For valuation practice, the results support benchmarking historical "
        "volatility against earlier option-implied information and reporting "
        "posterior model-price dispersion separately from the PI–PM correction. "
        "Future work should expand the number of valuation dates and liquid "
        "contracts and examine stochastic-volatility and multifactor alternatives.",
    )

    repair_equation_references(doc)


def format_tables(doc):
    # All table cells remain native and editable.
    for idx, table in enumerate(doc.tables):
        table.autofit = False
        if idx < 3:
            table.cell(0, 1).text = "N"
        if idx == 0:
            table.cell(3, 0).text = "Mean |PI − PM|"
        widths = (
            [Cm(8.3), Cm(2.5), Cm(4.6)] if idx == 0 else
            [Cm(6.8), Cm(1.8), Cm(2), Cm(2), Cm(2.4)] if idx in (1, 2) else
            [Cm(8.0), Cm(3.8), Cm(3.8)]
        )
        for col, width in zip(table.columns, widths):
            col.width = width
        for i, row in enumerate(table.rows):
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement("w:cantSplit"))
            if i == 0:
                props.append(OxmlElement("w:tblHeader"))
            for cell, width in zip(row.cells, widths):
                cell.width = width
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.keep_with_next = (i < len(table.rows) - 1)
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10)


def format_layout(doc):
    # Margins and fonts are from REMEF author instructions, not from the typeset PDF.
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
        sec.top_margin = sec.bottom_margin = Cm(2.5)
        sec.left_margin = sec.right_margin = Cm(3)
        # Page numbers are mandatory in REMEF submissions. Use a real Word
        # PAGE field rather than a hardcoded number that would not update.
        footer = sec.footer.paragraphs[0]
        footer.text = ""
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(0)
        footer.paragraph_format.space_after = Pt(0)
        run = footer.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        display_run = OxmlElement("w:r")
        display_text = OxmlElement("w:t")
        display_text.text = "1"
        display_run.append(display_text)
        field.append(display_run)
        footer._p.append(field)
    styles = {p.style.name: p.style for p in doc.paragraphs}
    for name in ("Normal", "Body Text", "First Paragraph", "Bibliography",
                 "Heading 1", "Heading 2", "Heading 3"):
        if name not in styles:
            continue
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.space_after = Pt(3)
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.space_before = Pt(9)
            style.paragraph_format.space_after = Pt(5)
            style.paragraph_format.keep_with_next = True

    for p in doc.paragraphs:
        if p.style.name in ("Body Text", "Normal", "First Paragraph"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in p.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
            if p.style.name.startswith("Heading"):
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.font.bold = True

    # Match the approved math-reference PDF's flush-left bilingual titles
    # while retaining REMEF's mandatory Times New Roman 12-pt Word typography.
    # Identify the titles by content, not by position: the unblinded document
    # includes author/contact paragraphs before its titles.
    publication_titles = {
        "Posterior Integration and Volatility Inputs in WTI Asian Options",
        "Integración posterior y volatilidad en opciones asiáticas WTI",
    }
    for p in doc.paragraphs:
        if p.text.strip() in publication_titles:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.font.bold = True

    # JEL codes explicitly labelled in both languages.
    english = find(doc, "Keywords:")
    spanish = find(doc, "JEL: C11;")
    insert = english.insert_paragraph_before("JEL classification: C11; C15; C58; G13; G17; Q40.")
    english._p.addnext(insert._p)
    spanish.text = "Clasificación JEL: C11; C15; C58; G13; G17; Q40."

    for p in list(doc.paragraphs):
        if p.text.startswith(("Figure ", "Table ")):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_before = Pt(4)
            if p.text.startswith("Figure ") and "Source:" in p.text:
                cap, source = p.text.split("Source:", 1)
                p.text = cap.strip()
                extra = p.insert_paragraph_before()
                p._p.addnext(extra._p)
                extra.text = "Source: " + source.strip()
                for run in extra.runs:
                    run.italic = True
                    run.font.size = Pt(10)
            if p.text.startswith("Table 2.") and " Panel A: main out-of-sample comparison" in p.text:
                p.text = p.text.replace(" Panel A: main out-of-sample comparison", "")
                panel = p.insert_paragraph_before("Panel A: Main out-of-sample comparison")
                p._p.addnext(panel._p)
                panel.alignment = WD_ALIGN_PARAGRAPH.LEFT
                panel.paragraph_format.keep_with_next = True
                for run in panel.runs:
                    run.font.bold = True
                    run.font.size = Pt(10)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
        if p.text.startswith("Panel B:"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            for run in p.runs:
                run.font.size = Pt(10)
                run.font.bold = True
        if p.text.startswith("Source:"):
            for run in p.runs:
                run.font.size = Pt(10)
                run.italic = True
        if p.style.name == "Bibliography":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)

    # The author has explicitly confirmed no external funding and no conflicts
    # of interest. Keep those statements in BOTH reviewer and editorial copies.
    # No unconfirmed acknowledgment text should be introduced here.
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""


def fit_figures_to_word_text_width(doc) -> None:
    """Use the actual Word section width rather than rescaling a wide plot.

    The REMEF manuscript is US Letter with 3 cm left and right margins.
    6.1375 inches is the physical text width; the Matplotlib figures are
    independently authored at this same width. Do not alter image pixels.
    """
    sec = doc.sections[0]
    target_width = int(sec.page_width - sec.left_margin - sec.right_margin)
    for shape in doc.inline_shapes:
        original_width, original_height = int(shape.width), int(shape.height)
        if original_width <= 0 or original_height <= 0:
            continue
        shape.width = target_width
        shape.height = int(round(original_height * target_width / original_width))


def prepare_docx(path: Path) -> None:
    doc = Document(path)
    adapt_sections(doc)
    format_layout(doc)
    format_tables(doc)
    fit_figures_to_word_text_width(doc)
    doc.save(path)


if __name__ == "__main__":
    import sys
    prepare_docx(Path(sys.argv[1]))
